from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path

import pytesseract
from openpyxl import Workbook
from openpyxl.styles import Font
from PIL import Image

from ..db import LIBRARY_DIR, connect, dumps, loads, new_id, utc_now


def text_to_fields(raw_text: str) -> dict[str, str]:
    """Turn OCR text into generic Field1..FieldN values.

    AEC title blocks vary a lot, so we do not assume labels like
    Drawing No / Rev. Each meaningful line (or table-like cell) becomes a field.
    """
    fields: dict[str, str] = {}
    if not raw_text or not raw_text.strip():
        return fields

    chunks: list[str] = []
    for line in raw_text.splitlines():
        line = line.strip()
        if not line:
            continue
        # Split dense table-like rows on 2+ spaces or pipes.
        parts = [p.strip() for p in re.split(r"\s{2,}|\s*\|\s*", line) if p.strip()]
        if len(parts) >= 2:
            chunks.extend(parts)
        else:
            chunks.append(line)

    # De-duplicate consecutive repeats from noisy OCR.
    cleaned: list[str] = []
    for chunk in chunks:
        chunk = re.sub(r"\s+", " ", chunk).strip(" .-:")
        if not chunk:
            continue
        if cleaned and cleaned[-1].lower() == chunk.lower():
            continue
        cleaned.append(chunk)

    for index, value in enumerate(cleaned, start=1):
        fields[f"Field{index}"] = value
    return fields


def ocr_info_block_layout(image_path: Path) -> tuple[str, dict[str, str], list[dict], str]:
    """OCR Info Block and keep approximate original layout boxes (0-1 in crop)."""
    with Image.open(image_path) as img:
        rgb = img.convert("RGB")
        width, height = rgb.size
        raw = pytesseract.image_to_string(rgb, config="--psm 6").strip()
        data = pytesseract.image_to_data(rgb, output_type=pytesseract.Output.DICT, config="--psm 6")

    words: list[dict] = []
    n = len(data.get("text") or [])
    for i in range(n):
        text = str(data["text"][i] or "").strip()
        if not text:
            continue
        try:
            conf = float(data["conf"][i])
        except Exception:
            conf = -1.0
        if conf < 0:
            continue
        words.append(
            {
                "text": text,
                "x": int(data["left"][i]),
                "y": int(data["top"][i]),
                "w": int(data["width"][i]),
                "h": int(data["height"][i]),
            }
        )

    # Cluster words into visual lines by vertical proximity (keeps title-block layout).
    words.sort(key=lambda w: (w["y"], w["x"]))
    line_groups: list[list[dict]] = []
    for word in words:
        cy = word["y"] + word["h"] / 2.0
        placed = False
        for group in line_groups:
            gy = sum(g["y"] + g["h"] / 2.0 for g in group) / len(group)
            gh = max(g["h"] for g in group)
            if abs(cy - gy) <= max(8.0, gh * 0.7):
                group.append(word)
                placed = True
                break
        if not placed:
            line_groups.append([word])

    layout: list[dict] = []
    fields: dict[str, str] = {}
    for group in line_groups:
        group.sort(key=lambda w: w["x"])
        value = re.sub(r"\s+", " ", " ".join(w["text"] for w in group)).strip(" .-:")
        if not value:
            continue
        if fields and list(fields.values())[-1].lower() == value.lower():
            continue
        x0 = min(w["x"] for w in group)
        y0 = min(w["y"] for w in group)
        x1 = max(w["x"] + w["w"] for w in group)
        y1 = max(w["y"] + w["h"] for w in group)
        key = f"Field{len(fields) + 1}"
        box = {
            "x": max(0.0, x0 / float(width)),
            "y": max(0.0, y0 / float(height)),
            "w": max(0.02, (x1 - x0) / float(width)),
            "h": max(0.02, (y1 - y0) / float(height)),
        }
        box["w"] = min(box["w"], 1.0 - box["x"])
        box["h"] = min(box["h"], 1.0 - box["y"])
        # Keep field height compact for editable overlays.
        box["h"] = min(box["h"], 0.12)
        fields[key] = value
        layout.append({"key": key, "text": value, "box": box})

    if not fields:
        fields = text_to_fields(raw)
        n = max(len(fields), 1)
        for index, (key, value) in enumerate(fields.items(), start=1):
            layout.append(
                {
                    "key": key,
                    "text": value,
                    "box": {
                        "x": 0.04,
                        "y": 0.04 + (index - 1) * (0.9 / n),
                        "w": 0.92,
                        "h": max(0.04, 0.85 / n),
                    },
                }
            )

    return raw or "\n".join(fields.values()), fields, layout, "ocr"


def ocr_info_block_image(image_path: Path) -> tuple[str, dict[str, str], str]:
    raw, fields, _layout, method = ocr_info_block_layout(image_path)
    return raw, fields, method


def _row_payload(row) -> dict:
    fields = loads(row["fields_json"])
    layout = loads(row["layout_json"]) if "layout_json" in row.keys() and row["layout_json"] else []
    return {
        "id": row["id"],
        "segment_id": row["segment_id"],
        "page_id": row["page_id"],
        "source_filename": row["source_filename"],
        "segment_filename": row["segment_filename"],
        "page_index": row["page_index"],
        "raw_text": row["raw_text"],
        "fields": fields,
        "layout": layout,
        "field_count": row["field_count"],
        "method": row["method"],
        "created_at": row["created_at"],
        "updated_at": row["updated_at"],
        "image_url": f"/media/library/Info Block/{row['segment_filename']}",
    }


def list_info_block_rows() -> list[dict]:
    with connect() as conn:
        rows = conn.execute(
            """
            SELECT * FROM info_block_rows
            ORDER BY updated_at DESC
            """
        ).fetchall()
    return [_row_payload(row) for row in rows]


def ocr_all_info_blocks() -> dict:
    """OCR every Info Block segment in the library into Field1..N rows."""
    with connect() as conn:
        segments = conn.execute(
            """
            SELECT s.*, p.page_index, d.filename AS source_filename
            FROM segments s
            JOIN pages p ON p.id = s.page_id
            JOIN drawings d ON d.id = p.drawing_id
            WHERE s.kind = 'info_block'
            ORDER BY s.created_at ASC
            """
        ).fetchall()

    processed = []
    errors = []
    now = utc_now()

    for seg in segments:
        image_path = LIBRARY_DIR / seg["relative_path"]
        if not image_path.exists():
            errors.append(
                {
                    "segment_id": seg["id"],
                    "error": f"Missing file: {seg['relative_path']}",
                }
            )
            continue

        try:
            raw_text, fields, layout, method = ocr_info_block_layout(image_path)
        except Exception as exc:  # noqa: BLE001 - collect and continue
            errors.append({"segment_id": seg["id"], "error": str(exc)})
            continue

        with connect() as conn:
            existing = conn.execute(
                "SELECT id FROM info_block_rows WHERE segment_id = ?",
                (seg["id"],),
            ).fetchone()
            if existing:
                conn.execute(
                    """
                    UPDATE info_block_rows
                    SET source_filename = ?, segment_filename = ?, page_index = ?,
                        raw_text = ?, fields_json = ?, field_count = ?, method = ?,
                        layout_json = ?, updated_at = ?
                    WHERE segment_id = ?
                    """,
                    (
                        seg["source_filename"],
                        seg["filename"],
                        seg["page_index"],
                        raw_text,
                        dumps(fields),
                        len(fields),
                        method,
                        dumps(layout),
                        now,
                        seg["id"],
                    ),
                )
                row_id = existing["id"]
            else:
                row_id = new_id("ibr")
                conn.execute(
                    """
                    INSERT INTO info_block_rows
                    (id, segment_id, page_id, source_filename, segment_filename,
                     page_index, raw_text, fields_json, field_count, method,
                     layout_json, created_at, updated_at)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        row_id,
                        seg["id"],
                        seg["page_id"],
                        seg["source_filename"],
                        seg["filename"],
                        seg["page_index"],
                        raw_text,
                        dumps(fields),
                        len(fields),
                        method,
                        dumps(layout),
                        now,
                        now,
                    ),
                )

        processed.append(
            {
                "id": row_id,
                "segment_id": seg["id"],
                "source_filename": seg["source_filename"],
                "segment_filename": seg["filename"],
                "field_count": len(fields),
                "fields": fields,
                "method": method,
            }
        )

    return {
        "processed": len(processed),
        "errors": errors,
        "rows": processed,
        "all_rows": list_info_block_rows(),
    }


def build_info_block_excel() -> bytes:
    rows = list_info_block_rows()
    max_fields = max((row["field_count"] for row in rows), default=0)

    wb = Workbook()
    ws = wb.active
    ws.title = "Info Block"

    headers = [
        "Source File",
        "Page",
        "Info Block JPG",
        "Method",
        "Field Count",
    ] + [f"Field{i}" for i in range(1, max_fields + 1)]

    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in rows:
        values = [
            row["source_filename"],
            row["page_index"] + 1,
            row["segment_filename"],
            row["method"],
            row["field_count"],
        ]
        for i in range(1, max_fields + 1):
            values.append(row["fields"].get(f"Field{i}", ""))
        ws.append(values)

    # Raw text sheet for review / debugging.
    raw_ws = wb.create_sheet("Raw Text")
    raw_ws.append(["Source File", "Page", "Info Block JPG", "Raw Text"])
    for cell in raw_ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        raw_ws.append(
            [
                row["source_filename"],
                row["page_index"] + 1,
                row["segment_filename"],
                row["raw_text"],
            ]
        )

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
